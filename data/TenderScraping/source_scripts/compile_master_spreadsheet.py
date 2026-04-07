#!/usr/bin/env python3
"""
Master Spreadsheet Compiler - Updated to include all expansion files
Combines all JSON files into a comprehensive Excel workbook
"""

import json
import glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# Load all JSON files
all_entries = []
for json_file in sorted(glob.glob('/sessions/peaceful-zen-gauss/tender_db/*.json')):
    with open(json_file, 'r') as f:
        entries = json.load(f)
        all_entries.extend(entries)
        print(f"Loaded {len(entries)} entries from {json_file.split('/')[-1]}")

print(f"\nTotal entries: {len(all_entries)}")

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Tender Sources"

# Define headers
headers = [
    "Source Name",
    "URL",
    "Country Code",
    "Country Name",
    "Region",
    "Entity Type",
    "Entity Subtype",
    "Sectors",
    "Language",
    "Update Frequency",
    "Registration Required",
    "Registration Type",
    "API Available",
    "RSS Feed",
    "Email Alerts",
    "Est. Annual Tenders",
    "Tender Value Range",
    "Primary Contact",
    "Technical Notes",
    "Last Verified",
    "Data Quality Score"
]

# Define styles
header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
header_fill = PatternFill("solid", fgColor="2E7D32")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

data_font = Font(name="Arial", size=9)
data_alignment = Alignment(vertical="top", wrap_text=True)
url_font = Font(name="Arial", size=9, color="0000FF", underline="single")

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Write headers
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# Freeze header row
ws.freeze_panes = "A2"

# Write data
for row_idx, entry in enumerate(all_entries, 2):
    ws.cell(row=row_idx, column=1, value=entry.get("source_name", "")).font = data_font
    ws.cell(row=row_idx, column=1).alignment = data_alignment
    ws.cell(row=row_idx, column=1).border = thin_border

    url_cell = ws.cell(row=row_idx, column=2, value=entry.get("url", ""))
    url_cell.font = url_font
    url_cell.alignment = data_alignment
    url_cell.border = thin_border

    ws.cell(row=row_idx, column=3, value=entry.get("country_code", "")).font = data_font
    ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=3).border = thin_border

    ws.cell(row=row_idx, column=4, value=entry.get("country_name", "")).font = data_font
    ws.cell(row=row_idx, column=4).alignment = data_alignment
    ws.cell(row=row_idx, column=4).border = thin_border

    ws.cell(row=row_idx, column=5, value=entry.get("region", "")).font = data_font
    ws.cell(row=row_idx, column=5).alignment = data_alignment
    ws.cell(row=row_idx, column=5).border = thin_border

    ws.cell(row=row_idx, column=6, value=entry.get("entity_type", "")).font = data_font
    ws.cell(row=row_idx, column=6).alignment = data_alignment
    ws.cell(row=row_idx, column=6).border = thin_border

    ws.cell(row=row_idx, column=7, value=entry.get("entity_subtype", "")).font = data_font
    ws.cell(row=row_idx, column=7).alignment = data_alignment
    ws.cell(row=row_idx, column=7).border = thin_border

    sectors = entry.get("sectors", [])
    sectors_str = ", ".join(sectors) if isinstance(sectors, list) else str(sectors)
    ws.cell(row=row_idx, column=8, value=sectors_str).font = data_font
    ws.cell(row=row_idx, column=8).alignment = data_alignment
    ws.cell(row=row_idx, column=8).border = thin_border

    ws.cell(row=row_idx, column=9, value=entry.get("language", "")).font = data_font
    ws.cell(row=row_idx, column=9).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=9).border = thin_border

    ws.cell(row=row_idx, column=10, value=entry.get("update_frequency", "")).font = data_font
    ws.cell(row=row_idx, column=10).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=10).border = thin_border

    reg_req = entry.get("registration_required")
    ws.cell(row=row_idx, column=11, value="Yes" if reg_req else "No" if reg_req is False else "").font = data_font
    ws.cell(row=row_idx, column=11).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=11).border = thin_border

    ws.cell(row=row_idx, column=12, value=entry.get("registration_type", "")).font = data_font
    ws.cell(row=row_idx, column=12).alignment = data_alignment
    ws.cell(row=row_idx, column=12).border = thin_border

    api = entry.get("api_available")
    ws.cell(row=row_idx, column=13, value="Yes" if api else "No" if api is False else "").font = data_font
    ws.cell(row=row_idx, column=13).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=13).border = thin_border

    rss = entry.get("rss_feed")
    ws.cell(row=row_idx, column=14, value="Yes" if rss else "No" if rss is False else "").font = data_font
    ws.cell(row=row_idx, column=14).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=14).border = thin_border

    email = entry.get("email_alerts")
    ws.cell(row=row_idx, column=15, value="Yes" if email else "No" if email is False else "").font = data_font
    ws.cell(row=row_idx, column=15).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=15).border = thin_border

    ws.cell(row=row_idx, column=16, value=entry.get("estimated_annual_tenders", "")).font = data_font
    ws.cell(row=row_idx, column=16).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=16).border = thin_border

    ws.cell(row=row_idx, column=17, value=entry.get("tender_value_range", "")).font = data_font
    ws.cell(row=row_idx, column=17).alignment = data_alignment
    ws.cell(row=row_idx, column=17).border = thin_border

    ws.cell(row=row_idx, column=18, value=entry.get("primary_contact", "")).font = data_font
    ws.cell(row=row_idx, column=18).alignment = data_alignment
    ws.cell(row=row_idx, column=18).border = thin_border

    ws.cell(row=row_idx, column=19, value=entry.get("technical_notes", "")).font = data_font
    ws.cell(row=row_idx, column=19).alignment = data_alignment
    ws.cell(row=row_idx, column=19).border = thin_border

    ws.cell(row=row_idx, column=20, value=entry.get("last_verified", "")).font = data_font
    ws.cell(row=row_idx, column=20).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=20).border = thin_border

    ws.cell(row=row_idx, column=21, value=entry.get("data_quality_score", "")).font = data_font
    ws.cell(row=row_idx, column=21).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=21).border = thin_border

# Set column widths
column_widths = {
    1: 45, 2: 55, 3: 12, 4: 20, 5: 25, 6: 22, 7: 25, 8: 40, 9: 12,
    10: 15, 11: 12, 12: 18, 13: 12, 14: 10, 15: 12, 16: 18, 17: 18,
    18: 25, 19: 45, 20: 14, 21: 15
}

for col_idx, width in column_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# Add auto-filter
ws.auto_filter.ref = ws.dimensions

# Create Summary sheet
summary_ws = wb.create_sheet("Summary")
summary_header_fill = PatternFill("solid", fgColor="1565C0")

summary_ws['A1'] = "Tender Intelligence Database - Summary"
summary_ws['A1'].font = Font(bold=True, size=16, name="Arial")
summary_ws.merge_cells('A1:D1')

summary_ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
summary_ws['A2'].font = Font(italic=True, size=10, name="Arial")

summary_ws['A3'] = f"Total Sources: {len(all_entries)}"
summary_ws['A3'].font = Font(bold=True, size=12, name="Arial")

# Count by country
country_counts = {}
for entry in all_entries:
    country = entry.get("country_name", "Unknown")
    country_counts[country] = country_counts.get(country, 0) + 1

# Count by entity type
entity_counts = {}
for entry in all_entries:
    entity = entry.get("entity_type", "Unknown")
    entity_counts[entity] = entity_counts.get(entity, 0) + 1

# Count by region
region_mapping = {
    "Africa": ["KEN", "TZA", "UGA", "GHA", "NGA", "ZAF", "ETH", "RWA", "ZMB", "ZWE", "BWA", "NAM", "MOZ", "MWI", "AGO", "COD", "CMR", "SEN", "CIV", "BFA", "MLI", "NER", "BEN", "TGO", "GIN", "SLE", "LBR", "GMB", "GNB", "MRT", "CPV", "SSD", "SDN", "ERI", "DJI", "SOM", "COM", "MUS", "SYC", "MDG", "REU", "GAB", "COG", "CAF", "TCD", "GNQ", "STP", "LSO", "SWZ", "BDI", "EGY", "MAR", "TUN", "DZA", "LBY"],
    "Latin America & Caribbean": ["BRA", "MEX", "ARG", "COL", "PER", "CHL", "ECU", "BOL", "PRY", "URY", "VEN", "GUY", "SUR", "GTM", "HND", "SLV", "NIC", "CRI", "PAN", "CUB", "DOM", "HTI", "JAM", "TTO", "BRB", "BHS", "BLZ"],
    "South Asia": ["IND", "PAK", "BGD", "LKA", "NPL", "AFG", "BTN", "MDV"],
    "Southeast Asia & Pacific": ["IDN", "PHL", "VNM", "THA", "MYS", "SGP", "MMR", "KHM", "LAO", "BRN", "TLS", "PNG", "FJI", "SLB", "VUT", "WSM", "TON", "KIR", "FSM", "MHL", "PLW", "NRU", "TUV", "COK", "NIU"],
    "Middle East & Central Asia": ["IRQ", "JOR", "LBN", "PSE", "YEM", "IRN", "ARE", "SAU", "QAT", "KWT", "BHR", "OMN", "KAZ", "UZB", "KGZ", "TJK", "TKM", "AZE", "ARM", "GEO", "MNG"],
    "International Organizations": ["INT", "Regional", "UN"]
}

region_counts = {k: 0 for k in region_mapping.keys()}
region_counts["Other"] = 0
for entry in all_entries:
    code = entry.get("country_code", "")
    found = False
    for region, codes in region_mapping.items():
        if code in codes:
            region_counts[region] += 1
            found = True
            break
    if not found:
        region_counts["Other"] += 1

# Write region summary
summary_ws['A5'] = "Sources by Region"
summary_ws['A5'].font = Font(bold=True, size=12, name="Arial")

row = 6
for col_idx, header in enumerate(["Region", "Count", "Percentage"], 1):
    cell = summary_ws.cell(row=row, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = summary_header_fill
    cell.border = thin_border

row = 7
for region, count in sorted(region_counts.items(), key=lambda x: -x[1]):
    summary_ws.cell(row=row, column=1, value=region).border = thin_border
    summary_ws.cell(row=row, column=2, value=count).border = thin_border
    summary_ws.cell(row=row, column=3, value=f"{count/len(all_entries)*100:.1f}%").border = thin_border
    row += 1

# Write entity type summary
row += 2
summary_ws.cell(row=row, column=1, value="Sources by Entity Type")
summary_ws.cell(row=row, column=1).font = Font(bold=True, size=12, name="Arial")

row += 1
for col_idx, header in enumerate(["Entity Type", "Count", "Percentage"], 1):
    cell = summary_ws.cell(row=row, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = summary_header_fill
    cell.border = thin_border

row += 1
for entity, count in sorted(entity_counts.items(), key=lambda x: -x[1])[:15]:
    summary_ws.cell(row=row, column=1, value=entity).border = thin_border
    summary_ws.cell(row=row, column=2, value=count).border = thin_border
    summary_ws.cell(row=row, column=3, value=f"{count/len(all_entries)*100:.1f}%").border = thin_border
    row += 1

# Write top countries
row += 2
summary_ws.cell(row=row, column=1, value="Top 25 Countries by Source Count")
summary_ws.cell(row=row, column=1).font = Font(bold=True, size=12, name="Arial")

row += 1
for col_idx, header in enumerate(["Country", "Count", "Percentage"], 1):
    cell = summary_ws.cell(row=row, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = summary_header_fill
    cell.border = thin_border

row += 1
for country, count in sorted(country_counts.items(), key=lambda x: -x[1])[:25]:
    summary_ws.cell(row=row, column=1, value=country).border = thin_border
    summary_ws.cell(row=row, column=2, value=count).border = thin_border
    summary_ws.cell(row=row, column=3, value=f"{count/len(all_entries)*100:.1f}%").border = thin_border
    row += 1

# Set column widths for summary
summary_ws.column_dimensions['A'].width = 35
summary_ws.column_dimensions['B'].width = 12
summary_ws.column_dimensions['C'].width = 12

# Save workbook
output_path = '/sessions/peaceful-zen-gauss/mnt/TenderScraping/Tender_Intelligence_Database.xlsx'
wb.save(output_path)
print(f"\nSaved workbook to: {output_path}")
