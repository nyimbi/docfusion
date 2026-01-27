#!/usr/bin/env python3
"""
Document Processor

Advanced document processing with multi-format support including PDF, DOCX, HTML,
and text extraction with structure preservation and encoding detection.
"""

import asyncio
import logging
import mimetypes
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import chardet

try:
    from uuid_extensions import uuid7str
except ImportError:
    from uuid import uuid4

    def uuid7str() -> str:
        return str(uuid4())


# PDF processing
try:
    import pdfplumber
    import PyPDF2

    HAS_PDF_SUPPORT = True
except ImportError:
    HAS_PDF_SUPPORT = False

# DOCX processing
try:
    from docx import Document as DocxDocument
    from docx.shared import Inches

    HAS_DOCX_SUPPORT = True
except ImportError:
    HAS_DOCX_SUPPORT = False

# HTML processing
try:
    import html2text
    from bs4 import BeautifulSoup

    HAS_HTML_SUPPORT = True
except ImportError:
    HAS_HTML_SUPPORT = False

# Image OCR support
try:
    import pytesseract
    from PIL import Image

    HAS_OCR_SUPPORT = True
except ImportError:
    HAS_OCR_SUPPORT = False


class DocumentMetadata:
    """Document metadata extracted during processing"""

    def __init__(self):
        self.title: Optional[str] = None
        self.author: Optional[str] = None
        self.subject: Optional[str] = None
        self.creator: Optional[str] = None
        self.producer: Optional[str] = None
        self.creation_date: Optional[datetime] = None
        self.modification_date: Optional[datetime] = None
        self.page_count: int = 0
        self.word_count: int = 0
        self.character_count: int = 0
        self.language: Optional[str] = None
        self.encoding: Optional[str] = None
        self.content_type: str = ""
        self.file_size: int = 0
        self.custom_properties: Dict[str, Any] = {}

    def to_dict(self) -> Dict[str, Any]:
        """Convert metadata to dictionary"""
        return {
            "title": self.title,
            "author": self.author,
            "subject": self.subject,
            "creator": self.creator,
            "producer": self.producer,
            "creation_date": self.creation_date.isoformat()
            if self.creation_date
            else None,
            "modification_date": self.modification_date.isoformat()
            if self.modification_date
            else None,
            "page_count": self.page_count,
            "word_count": self.word_count,
            "character_count": self.character_count,
            "language": self.language,
            "encoding": self.encoding,
            "content_type": self.content_type,
            "file_size": self.file_size,
            "custom_properties": self.custom_properties,
        }


class ProcessingResult:
    """Result of document processing"""

    def __init__(self):
        self.success: bool = False
        self.content: str = ""
        self.structured_content: Dict[str, Any] = {}
        self.metadata: DocumentMetadata = DocumentMetadata()
        self.errors: List[str] = []
        self.warnings: List[str] = []
        self.processing_time: float = 0.0
        self.processor_used: str = ""
        self.extracted_images: List[Dict[str, Any]] = []
        self.tables: List[Dict[str, Any]] = []
        self.links: List[Dict[str, str]] = []


class DocumentProcessor:
    """Advanced document processor with multi-format support"""

    def __init__(self, config: Optional[Dict[str, Any]] = None):
        self.config = config or self._get_default_config()
        self.logger = logging.getLogger(__name__)

        # Validate dependencies
        self._validate_dependencies()

        self.logger.info("Document processor initialized")

    def _get_default_config(self) -> Dict[str, Any]:
        """Get default configuration"""
        return {
            "max_file_size": 50 * 1024 * 1024,  # 50MB
            "ocr_enabled": HAS_OCR_SUPPORT,
            "ocr_languages": ["eng"],
            "preserve_formatting": True,
            "extract_tables": True,
            "extract_images": True,
            "extract_links": True,
            "encoding_detection": True,
            "language_detection": True,
            "pdf_fallback_to_plumber": True,
            "html_parser": "lxml",
            "text_extraction_timeout": 300,  # 5 minutes
            "max_pages_processed": 1000,
        }

    def _validate_dependencies(self):
        """Validate required dependencies"""
        missing_deps = []

        if not HAS_PDF_SUPPORT:
            missing_deps.append("PDF support (PyPDF2, pdfplumber)")

        if not HAS_DOCX_SUPPORT:
            missing_deps.append("DOCX support (python-docx)")

        if not HAS_HTML_SUPPORT:
            missing_deps.append("HTML support (beautifulsoup4, html2text)")

        if missing_deps:
            self.logger.warning(
                f"Missing optional dependencies: {', '.join(missing_deps)}"
            )

    async def process_document(
        self, file_path: Union[str, Path], content_type: Optional[str] = None
    ) -> ProcessingResult:
        """Process document from file path"""
        start_time = asyncio.get_event_loop().time()
        result = ProcessingResult()

        try:
            file_path = Path(file_path)

            # Validate file
            if not file_path.exists():
                result.errors.append("File does not exist")
                return result

            # Check file size
            file_size = file_path.stat().st_size
            if file_size > self.config["max_file_size"]:
                result.errors.append(
                    f"File size ({file_size} bytes) exceeds maximum ({self.config['max_file_size']} bytes)"
                )
                return result

            # Detect content type
            if not content_type:
                content_type, _ = mimetypes.guess_type(str(file_path))
                if not content_type:
                    content_type = self._detect_content_type_by_content(file_path)

            result.metadata.content_type = content_type or "application/octet-stream"
            result.metadata.file_size = file_size

            # Read file content
            with open(file_path, "rb") as f:
                file_content = f.read()

            # Process based on content type
            await self._process_by_content_type(file_content, result, content_type)

            # Post-processing
            await self._post_process_content(result)

            result.success = len(result.errors) == 0
            result.processing_time = asyncio.get_event_loop().time() - start_time

            self.logger.info(
                f"Processed document: {file_path.name}, success: {result.success}, time: {result.processing_time:.2f}s"
            )

        except Exception as e:
            result.errors.append(f"Processing error: {str(e)}")
            self.logger.error(f"Document processing failed: {e}")

        return result

    async def process_content(
        self, content: bytes, content_type: str, filename: Optional[str] = None
    ) -> ProcessingResult:
        """Process document from raw content"""
        start_time = asyncio.get_event_loop().time()
        result = ProcessingResult()

        try:
            # Validate content size
            if len(content) > self.config["max_file_size"]:
                result.errors.append(f"Content size exceeds maximum allowed")
                return result

            result.metadata.content_type = content_type
            result.metadata.file_size = len(content)

            # Process based on content type
            await self._process_by_content_type(content, result, content_type)

            # Post-processing
            await self._post_process_content(result)

            result.success = len(result.errors) == 0
            result.processing_time = asyncio.get_event_loop().time() - start_time

        except Exception as e:
            result.errors.append(f"Processing error: {str(e)}")
            self.logger.error(f"Content processing failed: {e}")

        return result

    async def _process_by_content_type(
        self, content: bytes, result: ProcessingResult, content_type: Optional[str]
    ):
        """Process content based on detected type"""
        if not content_type:
            result.errors.append("Unable to determine content type")
            return

        if content_type.startswith("application/pdf"):
            await self._process_pdf(content, result)

        elif content_type.startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml"
        ):
            await self._process_docx(content, result)

        elif content_type.startswith("text/html"):
            await self._process_html(content, result)

        elif content_type.startswith("text/"):
            await self._process_text(content, result)

        elif content_type.startswith("image/"):
            await self._process_image_ocr(content, result)

        else:
            # Try to process as text with encoding detection
            await self._process_unknown(content, result)

    async def _process_pdf(self, content: bytes, result: ProcessingResult):
        """Process PDF document"""
        if not HAS_PDF_SUPPORT:
            result.errors.append("PDF processing not available - missing dependencies")
            return

        result.processor_used = "PDF"

        try:
            import io

            # Try PyPDF2 first
            try:
                pdf_file = io.BytesIO(content)
                pdf_reader = PyPDF2.PdfReader(pdf_file)

                # Extract metadata
                if pdf_reader.metadata:
                    metadata = pdf_reader.metadata
                    result.metadata.title = metadata.get("/Title", "").strip()
                    result.metadata.author = metadata.get("/Author", "").strip()
                    result.metadata.subject = metadata.get("/Subject", "").strip()
                    result.metadata.creator = metadata.get("/Creator", "").strip()
                    result.metadata.producer = metadata.get("/Producer", "").strip()

                    # Parse dates
                    creation_date = metadata.get("/CreationDate")
                    if creation_date:
                        result.metadata.creation_date = self._parse_pdf_date(
                            creation_date
                        )

                result.metadata.page_count = len(pdf_reader.pages)

                # Extract text
                text_content = []
                for page_num, page in enumerate(pdf_reader.pages):
                    if page_num >= self.config["max_pages_processed"]:
                        result.warnings.append(
                            f"Stopped processing at page {page_num} (max limit reached)"
                        )
                        break

                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text_content.append(page_text)
                    except Exception as e:
                        result.warnings.append(
                            f"Failed to extract text from page {page_num + 1}: {e}"
                        )

                result.content = "\n\n".join(text_content)

            except Exception as e:
                if self.config["pdf_fallback_to_plumber"]:
                    # Fallback to pdfplumber
                    await self._process_pdf_with_plumber(content, result)
                else:
                    raise e

        except Exception as e:
            result.errors.append(f"PDF processing failed: {str(e)}")

    async def _process_pdf_with_plumber(self, content: bytes, result: ProcessingResult):
        """Process PDF with pdfplumber for better table/structure extraction"""
        try:
            import io

            pdf_file = io.BytesIO(content)

            with pdfplumber.open(pdf_file) as pdf:
                result.metadata.page_count = len(pdf.pages)

                # Extract metadata
                if hasattr(pdf, "metadata") and pdf.metadata:
                    metadata = pdf.metadata
                    result.metadata.title = metadata.get("Title", "").strip()
                    result.metadata.author = metadata.get("Author", "").strip()
                    result.metadata.subject = metadata.get("Subject", "").strip()
                    result.metadata.creator = metadata.get("Creator", "").strip()
                    result.metadata.producer = metadata.get("Producer", "").strip()

                text_content = []
                tables = []

                for page_num, page in enumerate(pdf.pages):
                    if page_num >= self.config["max_pages_processed"]:
                        break

                    try:
                        # Extract text
                        page_text = page.extract_text()
                        if page_text:
                            text_content.append(page_text)

                        # Extract tables if enabled
                        if self.config["extract_tables"]:
                            page_tables = page.extract_tables()
                            if page_tables:
                                for table_idx, table in enumerate(page_tables):
                                    tables.append(
                                        {
                                            "page": page_num + 1,
                                            "table_index": table_idx,
                                            "data": table,
                                            "row_count": len(table),
                                            "column_count": len(table[0])
                                            if table
                                            else 0,
                                        }
                                    )

                    except Exception as e:
                        result.warnings.append(
                            f"Failed to process page {page_num + 1}: {e}"
                        )

                result.content = "\n\n".join(text_content)
                result.tables = tables
                result.processor_used = "PDF (pdfplumber)"

        except Exception as e:
            result.errors.append(f"PDF processing with pdfplumber failed: {str(e)}")

    async def _process_docx(self, content: bytes, result: ProcessingResult):
        """Process DOCX document"""
        if not HAS_DOCX_SUPPORT:
            result.errors.append("DOCX processing not available - missing dependencies")
            return

        result.processor_used = "DOCX"

        try:
            import io

            docx_file = io.BytesIO(content)
            doc = DocxDocument(docx_file)

            # Extract metadata
            core_props = doc.core_properties
            result.metadata.title = core_props.title or ""
            result.metadata.author = core_props.author or ""
            result.metadata.subject = core_props.subject or ""
            result.metadata.creator = core_props.author or ""
            result.metadata.creation_date = core_props.created
            result.metadata.modification_date = core_props.modified

            # Extract text content
            text_content = []
            tables = []

            # Process paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)

            # Process tables if enabled
            if self.config["extract_tables"]:
                for table_idx, table in enumerate(doc.tables):
                    table_data = []
                    for row in table.rows:
                        row_data = []
                        for cell in row.cells:
                            row_data.append(cell.text.strip())
                        table_data.append(row_data)

                    tables.append(
                        {
                            "table_index": table_idx,
                            "data": table_data,
                            "row_count": len(table_data),
                            "column_count": len(table_data[0]) if table_data else 0,
                        }
                    )

            result.content = "\n\n".join(text_content)
            result.tables = tables
            result.metadata.page_count = 1  # DOCX doesn't have explicit pages like PDF

        except Exception as e:
            result.errors.append(f"DOCX processing failed: {str(e)}")

    async def _process_html(self, content: bytes, result: ProcessingResult):
        """Process HTML document"""
        if not HAS_HTML_SUPPORT:
            result.errors.append("HTML processing not available - missing dependencies")
            return

        result.processor_used = "HTML"

        try:
            # Detect encoding
            encoding = "utf-8"
            if self.config["encoding_detection"]:
                detected = chardet.detect(content)
                if detected["encoding"]:
                    encoding = detected["encoding"]

            result.metadata.encoding = encoding

            # Decode content
            try:
                html_content = content.decode(encoding)
            except UnicodeDecodeError:
                # Fallback to utf-8 with error handling
                html_content = content.decode("utf-8", errors="replace")
                result.warnings.append(
                    f"Encoding detection failed, used UTF-8 with error handling"
                )

            # Parse with BeautifulSoup
            soup = BeautifulSoup(html_content, self.config["html_parser"])

            # Extract metadata
            title_tag = soup.find("title")
            if title_tag:
                result.metadata.title = title_tag.get_text().strip()

            # Extract author from meta tags
            author_meta = soup.find("meta", attrs={"name": "author"})
            if author_meta:
                result.metadata.author = author_meta.get("content", "").strip()

            # Extract description
            desc_meta = soup.find("meta", attrs={"name": "description"})
            if desc_meta:
                result.metadata.subject = desc_meta.get("content", "").strip()

            # Extract links if enabled
            links = []
            if self.config["extract_links"]:
                for link in soup.find_all("a", href=True):
                    links.append(
                        {
                            "text": link.get_text().strip(),
                            "url": link["href"],
                            "title": link.get("title", ""),
                        }
                    )

            # Extract tables if enabled
            tables = []
            if self.config["extract_tables"]:
                for table_idx, table in enumerate(soup.find_all("table")):
                    table_data = []
                    rows = table.find_all("tr")
                    for row in rows:
                        cells = row.find_all(["td", "th"])
                        row_data = [cell.get_text().strip() for cell in cells]
                        if row_data:
                            table_data.append(row_data)

                    if table_data:
                        tables.append(
                            {
                                "table_index": table_idx,
                                "data": table_data,
                                "row_count": len(table_data),
                                "column_count": len(table_data[0]) if table_data else 0,
                            }
                        )

            # Convert to plain text
            h = html2text.HTML2Text()
            h.ignore_links = not self.config["extract_links"]
            h.ignore_images = not self.config["extract_images"]
            h.body_width = 0  # Don't wrap lines

            result.content = h.handle(html_content)
            result.links = links
            result.tables = tables

        except Exception as e:
            result.errors.append(f"HTML processing failed: {str(e)}")

    async def _process_text(self, content: bytes, result: ProcessingResult):
        """Process plain text document"""
        result.processor_used = "Text"

        try:
            # Detect encoding
            encoding = "utf-8"
            if self.config["encoding_detection"]:
                detected = chardet.detect(content)
                if detected["encoding"] and detected["confidence"] > 0.7:
                    encoding = detected["encoding"]

            result.metadata.encoding = encoding

            # Decode content
            try:
                text_content = content.decode(encoding)
            except UnicodeDecodeError:
                # Try common encodings
                for fallback_encoding in ["latin-1", "cp1252", "iso-8859-1"]:
                    try:
                        text_content = content.decode(fallback_encoding)
                        result.metadata.encoding = fallback_encoding
                        result.warnings.append(
                            f"Encoding detection failed, used {fallback_encoding}"
                        )
                        break
                    except UnicodeDecodeError:
                        continue
                else:
                    # Last resort - utf-8 with error handling
                    text_content = content.decode("utf-8", errors="replace")
                    result.warnings.append(
                        "All encoding attempts failed, used UTF-8 with error handling"
                    )

            result.content = text_content

        except Exception as e:
            result.errors.append(f"Text processing failed: {str(e)}")

    async def _process_image_ocr(self, content: bytes, result: ProcessingResult):
        """Process image using OCR"""
        if not HAS_OCR_SUPPORT or not self.config["ocr_enabled"]:
            result.errors.append("OCR processing not available")
            return

        result.processor_used = "OCR"

        try:
            import io

            # Open image
            image_file = io.BytesIO(content)
            image = Image.open(image_file)

            # Configure OCR
            ocr_config = "--oem 3 --psm 6"  # Default OCR configuration
            languages = "+".join(self.config["ocr_languages"])

            # Extract text
            extracted_text = pytesseract.image_to_string(
                image, lang=languages, config=ocr_config
            )

            result.content = extracted_text
            result.metadata.custom_properties["ocr_languages"] = self.config[
                "ocr_languages"
            ]
            result.metadata.custom_properties["image_format"] = image.format
            result.metadata.custom_properties["image_size"] = image.size

        except Exception as e:
            result.errors.append(f"OCR processing failed: {str(e)}")

    async def _process_unknown(self, content: bytes, result: ProcessingResult):
        """Process unknown content type by attempting text extraction"""
        result.processor_used = "Unknown/Text fallback"

        # Try to process as text
        await self._process_text(content, result)

        if not result.content.strip():
            result.warnings.append("No readable text content found")

    async def _post_process_content(self, result: ProcessingResult):
        """Post-process extracted content"""
        if not result.content:
            return

        # Calculate statistics
        result.metadata.character_count = len(result.content)
        result.metadata.word_count = len(result.content.split())

        # Language detection (simplified)
        if self.config["language_detection"]:
            result.metadata.language = self._detect_language(result.content)

        # Create structured content
        result.structured_content = {
            "title": result.metadata.title or "",
            "content": result.content,
            "paragraphs": self._extract_paragraphs(result.content),
            "sentences": self._extract_sentences(result.content),
            "tables": result.tables,
            "links": result.links,
            "statistics": {
                "word_count": result.metadata.word_count,
                "character_count": result.metadata.character_count,
                "paragraph_count": len(self._extract_paragraphs(result.content)),
                "sentence_count": len(self._extract_sentences(result.content)),
            },
        }

    def _detect_content_type_by_content(self, file_path: Path) -> Optional[str]:
        """Detect content type by examining file content"""
        try:
            with open(file_path, "rb") as f:
                header = f.read(1024)

            # PDF
            if header.startswith(b"%PDF"):
                return "application/pdf"

            # ZIP-based formats (DOCX, etc.)
            if header.startswith(b"PK\x03\x04"):
                # Could be DOCX, XLSX, etc.
                if file_path.suffix.lower() == ".docx":
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

            # HTML
            if b"<html" in header.lower() or b"<!doctype html" in header.lower():
                return "text/html"

            # Try to decode as text
            try:
                header.decode("utf-8")
                return "text/plain"
            except UnicodeDecodeError:
                pass

            return None

        except Exception:
            return None

    def _parse_pdf_date(self, date_string: str) -> Optional[datetime]:
        """Parse PDF date format"""
        try:
            # PDF date format: D:YYYYMMDDHHmmSSOHH'mm'
            if date_string.startswith("D:"):
                date_string = date_string[2:]

            # Extract basic date parts
            if len(date_string) >= 8:
                year = int(date_string[:4])
                month = int(date_string[4:6])
                day = int(date_string[6:8])
                hour = int(date_string[8:10]) if len(date_string) >= 10 else 0
                minute = int(date_string[10:12]) if len(date_string) >= 12 else 0
                second = int(date_string[12:14]) if len(date_string) >= 14 else 0

                return datetime(year, month, day, hour, minute, second)

        except (ValueError, IndexError):
            pass

        return None

    def _detect_language(self, text: str) -> Optional[str]:
        """Simple language detection (placeholder for more sophisticated detection)"""
        # This is a simplified implementation
        # In production, you might use langdetect or similar libraries

        if not text or len(text.split()) < 10:
            return None

        # Simple heuristics based on common words
        english_words = [
            "the",
            "and",
            "or",
            "but",
            "in",
            "on",
            "at",
            "to",
            "for",
            "of",
            "with",
            "by",
        ]
        text_lower = text.lower()

        english_count = sum(1 for word in english_words if word in text_lower)
        if english_count >= 3:
            return "en"

        return "unknown"

    def _extract_paragraphs(self, text: str) -> List[str]:
        """Extract paragraphs from text"""
        if not text:
            return []

        # Split by double newlines or more
        paragraphs = re.split(r"\n\s*\n", text.strip())
        return [p.strip() for p in paragraphs if p.strip()]

    def _extract_sentences(self, text: str) -> List[str]:
        """Extract sentences from text (simple implementation)"""
        if not text:
            return []

        # Simple sentence splitting
        sentences = re.split(r"[.!?]+", text)
        return [s.strip() for s in sentences if s.strip()]

    def get_supported_formats(self) -> List[str]:
        """Get list of supported file formats"""
        formats = ["text/plain"]

        if HAS_PDF_SUPPORT:
            formats.append("application/pdf")

        if HAS_DOCX_SUPPORT:
            formats.append(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

        if HAS_HTML_SUPPORT:
            formats.extend(["text/html", "application/xhtml+xml"])

        if HAS_OCR_SUPPORT and self.config["ocr_enabled"]:
            formats.extend(["image/jpeg", "image/png", "image/tiff", "image/bmp"])

        return formats

    def get_processor_info(self) -> Dict[str, Any]:
        """Get processor information and capabilities"""
        return {
            "supported_formats": self.get_supported_formats(),
            "dependencies": {
                "pdf_support": HAS_PDF_SUPPORT,
                "docx_support": HAS_DOCX_SUPPORT,
                "html_support": HAS_HTML_SUPPORT,
                "ocr_support": HAS_OCR_SUPPORT,
            },
            "config": self.config.copy(),
            "version": "1.0.0",
        }


# Factory function
def create_document_processor(
    config: Optional[Dict[str, Any]] = None,
) -> DocumentProcessor:
    """Create DocumentProcessor instance with configuration"""
    return DocumentProcessor(config)
